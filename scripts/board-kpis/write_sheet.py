#!/usr/bin/env python3
"""
Copy the reported months from the review workbook into the live KPI sheet.

Reads `.board-kpis/kpi_draft.xlsx` (built by draft_workbook.py and reviewed)
and the live workbook through the Sheets API, and plans a cell-by-cell write:

- the main tab and CAC Inputs: every non-empty cell in the reported months'
  columns (AM onward), formulas as formulas, with July's formatting copied
  across;
- Churn Inputs: the cancellations draft_workbook.py appended, added under the
  live sheet's last row;
- a Board Notes tab: the draft's Board Notes rows as a section for this
  report, and a note on every rebased cell.

July and earlier are never written. A live formula that differs from the
draft (the sheet carries some formulas ahead into future months, such as the
13-month TTM windows and the roll-forward Beginning rows) is replaced, and
the plan lists each one with the formula it replaces. A live typed value
that differs is a conflict: the plan lists it and nothing is written until it
is resolved (someone typed into the sheet after the export).

The plan is always written to `.board-kpis/sheet_plan.md`. Nothing is sent
unless BOARD_KPIS_WRITE=1. After a write, the main tab's calculated values
for the reported months are read back and compared with the draft as
LibreOffice calculates it.

The working-paper tabs (Sources, ledgers, CAC Variance, Ending MRR Recon, CAC
Excluded) are not written to the live sheet: draft_workbook.py saves them as
`.board-kpis/kpi_support.xlsx`, filed in Drive beside it. Put its link in
config.json `support_links` under the report label (e.g. "Report: August
2026 to September 2026") and Board Notes links to it.

Environment:
    WINS_SA_KEY        the service account JSON (see scripts/common/sheets.py);
                       the workbook is shared with it as an Editor
    BOARD_KPIS_WRITE   1 to write; anything else is a dry run

Usage:
    python3 write_sheet.py
"""

import datetime as dt
import json
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.utils import column_index_from_string as col_index
from openpyxl.utils import get_column_letter as col_letter

import draft_workbook as D
import workspace as w

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "common"))
from sheets import Sheet  # noqa: E402

PLAN = "sheet_plan.md"
NOTES_TAB = "Board Notes"
COLUMN_TABS = (D.MAIN, D.CAC)
CHURN_COLS = 6  # Client, Type, MRR, Start, End, Days
EPOCH = dt.datetime(1899, 12, 30)


def quote(tab):
    return "'%s'" % tab.replace("'", "''")


def cell_value(v):
    """A draft cell as the Sheets API takes it with USER_ENTERED."""
    if isinstance(v, dt.datetime):
        return (v - EPOCH).days
    return v


def same(live, draft):
    if isinstance(live, (int, float)) and isinstance(draft, (int, float)):
        return abs(live - draft) < 0.005
    return str(live).strip() == str(draft).strip()


def months_in(draft_wb):
    """Reported months: the main tab's columns after July that the draft filled."""
    ws = draft_wb[D.MAIN]
    return [c for c in range(D.TEMPLATE_COL + 1, ws.max_column + 1)
            if any(ws.cell(r, c).value is not None for r in range(6, ws.max_row + 1))]


def read(sheet, ranges, render="FORMULA"):
    out = sheet.call("GET", "/values:batchGet", params={"ranges": ranges, "valueRenderOption": render})
    return [v.get("values", []) for v in out["valueRanges"]]


def grid_get(rows, r, c):
    """rows from the API (0-based, ragged) -> value at 0-based r, c or None."""
    if r < len(rows) and c < len(rows[r]) and rows[r][c] != "":
        return rows[r][c]
    return None


def plan_columns(sheet, draft_wb, cols):
    writes, replaced, conflicts = [], [], []
    width = {t: p["gridProperties"]["columnCount"] for t, p in sheet_ids(sheet).items()}
    for tab in COLUMN_TABS:
        ws = draft_wb[tab]
        top = min(cols[-1], width[tab])  # the live tab may not reach the new months yet
        live = (read(sheet, ["%s!%s1:%s%d" % (quote(tab), col_letter(cols[0]), col_letter(top), ws.max_row)])[0]
                if top >= cols[0] else [])
        for r in range(1, ws.max_row + 1):
            for i, c in enumerate(cols):
                want = cell_value(ws.cell(r, c).value)
                if want is None:
                    continue
                have = grid_get(live, r - 1, i)
                ref = "%s%d" % (col_letter(c), r)
                if have is None:
                    writes.append((tab, ref, want))
                elif same(have, want):
                    continue
                elif isinstance(have, str) and have.startswith("="):
                    writes.append((tab, ref, want))
                    replaced.append((tab, ref, have, want))
                else:
                    conflicts.append((tab, ref, have, want))
    return writes, replaced, conflicts


def churn_rows(draft_wb):
    """The rows draft_workbook.py appended to Churn Inputs, as [client, type, mrr, start, end]."""
    export = openpyxl.load_workbook(w.path(D.EXPORT))[D.CHURN]
    last = max(i for i, r in enumerate(export.iter_rows(values_only=True), 1) if r[0])
    ws = draft_wb[D.CHURN]
    return [[cell_value(ws.cell(r, c).value) for c in range(1, CHURN_COLS)]
            for r in range(last + 1, ws.max_row + 1) if ws.cell(r, 1).value]


def plan_churn(sheet, rows):
    live = read(sheet, ["%s!A1:%s" % (quote(D.CHURN), col_letter(CHURN_COLS))])[0]
    last = max(i for i, r in enumerate(live, 1) if r and r[0])
    seen = {(str(r[0]).strip(), str(r[4]) if len(r) > 4 else "") for r in live if r}
    new = [r for r in rows if (str(r[0]).strip(), str(r[4])) not in seen]
    out = []
    for i, r in enumerate(new, last + 1):
        days = "=E%d-D%d" % (i, i) if r[3] and r[4] else None
        out.append((i, r + [days]))
    return last, out, len(rows) - len(new)


def rebase_notes(sources):
    """[(main tab cell, note)] for every rebased cell on the Sources tab."""
    out = []
    for month, ref, metric, value, source, detail in sources:
        if "(rebase)" in str(source) and "!" not in str(ref):
            out.append((ref, "Rebase, %s: %s. %s" % (month, source.replace(" (rebase)", ""), detail)))
    return out


def board_rows(draft_wb, label, support_link=None):
    ws = draft_wb[NOTES_TAB]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    link = [["Working papers (sources, ledgers, CAC variance, reconciliation)", support_link]] if support_link else []
    return [[label]] + link + rows + [[]]


def markdown(cols, writes, replaced, conflicts, churn, notes, board, skipped):
    lines = ["# Sheet write plan", "",
             "Columns: %s. Cells to write: %d. Conflicts: %d." % (
                 ", ".join(col_letter(c) for c in cols), len(writes), len(conflicts)), ""]
    if conflicts:
        lines += ["## Conflicts (nothing is written until these are resolved)", "",
                  "| Tab | Cell | Live | Draft |", "| --- | --- | --- | --- |"]
        lines += ["| %s | %s | %s | %s |" % c for c in conflicts]
        lines.append("")
    if replaced:
        lines += ["## Live formulas replaced", "",
                  "| Tab | Cell | Live | Draft |", "| --- | --- | --- | --- |"]
        lines += ["| %s | %s | `%s` | `%s` |" % c for c in replaced]
        lines.append("")
    lines += ["## Cells", "", "| Tab | Cell | Value |", "| --- | --- | --- |"]
    lines += ["| %s | %s | `%s` |" % (t, ref, v) for t, ref, v in writes]
    lines += ["", "## Churn Inputs rows", ""]
    lines += ["- row %d: %s" % (i, r) for i, r in churn[1]]
    if skipped:
        lines.append("- %d already on the live tab, skipped" % skipped)
    lines += ["", "## Cell notes", ""] + ["- %s: %s" % n for n in notes]
    lines += ["", "## Board Notes tab", "", "%d rows, starting with: %s" % (len(board), board[0][0])]
    return "\n".join(lines)


def sheet_ids(sheet):
    meta = sheet.call("GET", "", params={"fields": "sheets(properties(sheetId,title,gridProperties),merges)"})
    return {s["properties"]["title"]: dict(s["properties"], merges=s.get("merges", [])) for s in meta["sheets"]}


def unmerged_spans(merges, rows, first_col, last_col):
    """0-based [start, end) row spans under `rows` with no merge touching the
    columns first_col..last_col (0-based): the section header bands are merged
    across the months, and a format paste cannot cut through a merge."""
    blocked = set()
    for m in merges:
        if m["startColumnIndex"] <= last_col and m["endColumnIndex"] > first_col - 1:
            blocked.update(range(m["startRowIndex"], m["endRowIndex"]))
    spans, start = [], None
    for r in range(rows + 1):
        if r < rows and r not in blocked:
            start = r if start is None else start
        elif start is not None:
            spans.append((start, r))
            start = None
    return spans


def history(sheet, max_row):
    """The main tab and CAC Inputs through July, as formulas."""
    last = col_letter(D.TEMPLATE_COL)
    tabs = (D.MAIN, D.CAC)
    got = read(sheet, ["%s!A1:%s%d" % (quote(t), last, max_row) for t in tabs])
    return {t: g for t, g in zip(tabs, got)}


def changed_history(before, after):
    """Cells through July that differ between two history() reads. Appending
    columns to CAC Inputs makes Sheets shift references that point past the
    tab's last column (row 44's published formulas do), so this is checked
    after every write rather than assumed."""
    out = []
    for tab, rows in before.items():
        for r in range(max(len(rows), len(after[tab]))):
            for c in range(D.TEMPLATE_COL):
                a, b = grid_get(rows, r, c), grid_get(after[tab], r, c)
                if a != b:
                    out.append((tab, "%s%d" % (col_letter(c + 1), r + 1), a, b))
    return out


def write(sheet, cols, writes, churn, notes, board):
    props = sheet_ids(sheet)
    requests = []
    # Room for the new months, then July's formats across.
    for tab in COLUMN_TABS:
        p = props[tab]
        short = cols[-1] - p["gridProperties"]["columnCount"]
        if short > 0:
            requests.append({"appendDimension": {"sheetId": p["sheetId"], "dimension": "COLUMNS", "length": short}})
        rows = max(int(ref.lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ")) for t, ref, v in writes if t == tab) if any(
            t == tab for t, _, _ in writes) else 0
        for start, end in unmerged_spans(p["merges"], rows, D.TEMPLATE_COL - 1, cols[-1] - 1):
            requests.append({"copyPaste": {
                "source": {"sheetId": p["sheetId"], "startRowIndex": start, "endRowIndex": end,
                           "startColumnIndex": D.TEMPLATE_COL - 1, "endColumnIndex": D.TEMPLATE_COL},
                "destination": {"sheetId": p["sheetId"], "startRowIndex": start, "endRowIndex": end,
                                "startColumnIndex": cols[0] - 1, "endColumnIndex": cols[-1]},
                "pasteType": "PASTE_FORMAT"}})
    last, rows, _ = churn
    if rows:
        p = props[D.CHURN]
        requests.append({"copyPaste": {
            "source": {"sheetId": p["sheetId"], "startRowIndex": last - 1, "endRowIndex": last,
                       "startColumnIndex": 0, "endColumnIndex": CHURN_COLS},
            "destination": {"sheetId": p["sheetId"], "startRowIndex": last, "endRowIndex": last + len(rows),
                            "startColumnIndex": 0, "endColumnIndex": CHURN_COLS},
            "pasteType": "PASTE_FORMAT"}})
    if NOTES_TAB not in props:
        requests.append({"addSheet": {"properties": {"title": NOTES_TAB}}})
    if requests:
        sheet.call("POST", ":batchUpdate", json={"requests": requests})

    data = [{"range": "%s!%s" % (quote(t), ref), "values": [[v]]} for t, ref, v in writes]
    data += [{"range": "%s!A%d:%s%d" % (quote(D.CHURN), i, col_letter(CHURN_COLS), i), "values": [r]}
             for i, r in rows]
    existing = read(sheet, ["%s!A1:A" % quote(NOTES_TAB)])[0]
    start = len(existing) + (2 if existing else 1)
    data.append({"range": "%s!A%d" % (quote(NOTES_TAB), start),
                 "values": [[("" if v is None else v) for v in r] for r in board]})
    sheet.call("POST", "/values:batchUpdate", json={"valueInputOption": "USER_ENTERED", "data": data})

    main = sheet_ids(sheet)[D.MAIN]["sheetId"]
    note_requests = []
    for ref, text in notes:
        c = col_index(ref.rstrip("0123456789"))
        r = int(ref[len(ref.rstrip("0123456789")):])
        note_requests.append({"updateCells": {
            "range": {"sheetId": main, "startRowIndex": r - 1, "endRowIndex": r,
                      "startColumnIndex": c - 1, "endColumnIndex": c},
            "rows": [{"values": [{"note": text}]}], "fields": "note"}})
    if note_requests:
        sheet.call("POST", ":batchUpdate", json={"requests": note_requests})


def recalculated(path):
    """The draft as LibreOffice calculates it, or None without soffice."""
    if not shutil.which("soffice"):
        return None
    out = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, path],
                   capture_output=True, timeout=300)
    done = os.path.join(out, os.path.basename(path))
    return openpyxl.load_workbook(done, data_only=True) if os.path.exists(done) else None


def verify(sheet, cols, max_row):
    calc = recalculated(w.path(D.OUT))
    if calc is None:
        print("soffice not found: skipped the read-back comparison")
        return 0
    ws = calc[D.MAIN]
    live = read(sheet, ["%s!%s1:%s%d" % (quote(D.MAIN), col_letter(cols[0]), col_letter(cols[-1]), max_row)],
                render="UNFORMATTED_VALUE")[0]
    diffs = []
    for r in range(6, max_row + 1):
        for i, c in enumerate(cols):
            want, have = ws.cell(r, c).value, grid_get(live, r - 1, i)
            if isinstance(want, (int, float)) and not isinstance(want, bool):
                if not isinstance(have, (int, float)) or abs(have - want) > max(0.01, abs(want) * 1e-6):
                    diffs.append("%s%d: sheet %r, draft %r" % (col_letter(c), r, have, want))
    for d in diffs:
        print("  differs " + d)
    print("read-back: %d calculated cells differ from the draft" % len(diffs))
    return len(diffs)


def main():
    config = json.load(open(os.path.join(HERE, "config.json")))
    sheet = Sheet(os.environ["WINS_SA_KEY"], config["kpi_spreadsheet_id"])
    draft_wb = openpyxl.load_workbook(w.path(D.OUT))
    cols = months_in(draft_wb)
    if not cols or cols[0] <= D.TEMPLATE_COL:
        sys.exit("nothing to write after column %s" % col_letter(D.TEMPLATE_COL))

    writes, replaced, conflicts = plan_columns(sheet, draft_wb, cols)
    churn = plan_churn(sheet, churn_rows(draft_wb))
    sources = [[c.value for c in r] for r in draft_wb["Sources"].iter_rows(min_row=2) if r[1].value]
    notes = rebase_notes(sources)
    names = [dt.date(2023 + (c - 2 + 6) // 12, (c - 2 + 6) % 12 + 1, 1).strftime("%B %Y") for c in cols]
    label = "Report: %s, written %s" % (" to ".join(dict.fromkeys([names[0], names[-1]])), dt.date.today())
    board = board_rows(draft_wb, label, config.get("support_links", {}).get(label.split(",")[0]))
    last, rows, skipped = churn

    with open(w.path(PLAN), "w") as f:
        f.write(markdown(cols, writes, replaced, conflicts, churn, notes, board, skipped))
    print("%d cells (%d replace a live formula), %d Churn Inputs rows, %d cell notes, %d Board Notes rows, "
          "%d conflicts -> %s" % (len(writes), len(replaced), len(rows), len(notes), len(board),
                                   len(conflicts), w.path(PLAN)))
    if conflicts:
        for c in conflicts:
            print("  conflict %s!%s: live %r, draft %r" % c)
        return 1
    if os.environ.get("BOARD_KPIS_WRITE") != "1":
        print("dry run: set BOARD_KPIS_WRITE=1 to write")
        return 0
    rows = draft_wb[D.MAIN].max_row
    before = history(sheet, rows)
    write(sheet, cols, writes, churn, notes, board)
    print("written")
    moved = changed_history(before, history(sheet, rows))
    for tab, ref, a, b in moved:
        print("  history changed %s!%s: %r -> %r" % (tab, ref, a, b))
    if moved:
        print("%d cells through July changed as a side effect: restore them by hand or approve a restore" % len(moved))
    return 1 if verify(sheet, cols, rows) or moved else 0


if __name__ == "__main__":
    sys.exit(main())
